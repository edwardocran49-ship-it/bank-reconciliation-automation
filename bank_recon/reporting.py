"""Excel export for the reconciliation result."""

from __future__ import annotations

from collections import Counter, defaultdict
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .engine import ReconciliationResult


NAVY = "1F4E78"
LIGHT_BLUE = "D9EAF7"
WHITE = "FFFFFF"
DARK = "1F1F1F"
GREY = "595959"
GREEN = "008000"
RED = "C00000"
LIGHT_RED = "FCE8E6"
ACCOUNTING = '#,##0.00;(#,##0.00);-'
DATE_FORMAT = "yyyy-mm-dd"
THIN_GREY = Side(style="thin", color="808080")
DOUBLE_DARK = Side(style="double", color="1F1F1F")


def _excel_value(value: object) -> object:
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def _period_details(result: ReconciliationResult) -> tuple[str, str]:
    dates = pd.concat([result.bank["date"], result.gl["date"]])
    start, end = dates.min(), dates.max()
    if start.year == end.year and start.month == end.month:
        return end.strftime("%B %Y"), end.strftime("%b%Y").lower()
    return f"{start:%b %Y} to {end:%b %Y}", end.strftime("%b%Y").lower()


def _apply_base_style(sheet) -> None:
    sheet.sheet_view.showGridLines = False
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name="Aptos", size=11, color=DARK)
            cell.alignment = Alignment(vertical="center")


def _style_header(sheet, row: int, first_col: int, last_col: int) -> None:
    for column in range(first_col, last_col + 1):
        target = sheet.cell(row, column)
        target.fill = PatternFill("solid", fgColor=NAVY)
        target.font = Font(name="Aptos", size=11, bold=True, color=WHITE)
        target.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[row].height = 21


def _style_total(sheet, row: int, last_col: int) -> None:
    for cell in sheet[row][:last_col]:
        cell.font = Font(name="Aptos", size=11, bold=True, color=DARK)
        cell.border = Border(top=THIN_GREY, bottom=DOUBLE_DARK)


def _set_widths(sheet, widths: list[float]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_source_sheet(
    sheet,
    headers: list[str],
    rows: list[list[object]],
    table_name: str,
    widths: list[float],
    date_columns: tuple[int, ...],
    amount_columns: tuple[int, ...],
) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append([_excel_value(value) for value in row])
    _apply_base_style(sheet)
    sheet.freeze_panes = "A2"
    _style_header(sheet, 1, 1, len(headers))
    for row_index in range(2, sheet.max_row + 1):
        if row_index % 2 == 0:
            for cell in sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        for column in date_columns:
            sheet.cell(row_index, column).number_format = DATE_FORMAT
        for column in amount_columns:
            sheet.cell(row_index, column).number_format = ACCOUNTING
            sheet.cell(row_index, column).alignment = Alignment(
                horizontal="right", vertical="center"
            )
    _set_widths(sheet, widths)
    table = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(headers))}{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _source_views(result: ReconciliationResult) -> tuple[list[list[object]], list[list[object]]]:
    bank_matches = {
        row.bank_row_id: row for row in result.matches.itertuples(index=False)
    }
    gl_matches = {row.gl_row_id: row for row in result.matches.itertuples(index=False)}
    bank_exception_ids = set(
        result.exceptions.loc[result.exceptions["source"] == "BANK", "row_id"]
    )
    gl_exceptions = {
        row.row_id: row
        for row in result.exceptions[result.exceptions["source"] == "GL"].itertuples(index=False)
    }
    gl_reference_counts = Counter(result.gl["normalised_reference"])

    bank_rows: list[list[object]] = []
    for row in result.bank.itertuples(index=False):
        match = bank_matches.get(row.row_id)
        if match is not None:
            difference = round(float(match.bank_amount - match.gl_amount), 2)
            status = "Amount difference" if abs(difference) >= 0.005 else "Matched"
            gl_amount = float(match.gl_amount)
        else:
            difference = None
            gl_amount = None
            status = "Bank only" if row.row_id in bank_exception_ids else "Review"
        bank_rows.append(
            [row.date, row.description, row.reference, float(row.amount),
             gl_reference_counts.get(row.normalised_reference, 0), gl_amount, difference, status]
        )

    occurrence = defaultdict(int)
    gl_rows: list[list[object]] = []
    for row in result.gl.itertuples(index=False):
        match = gl_matches.get(row.row_id)
        key = (row.normalised_reference, round(float(row.amount), 2))
        duplicate_occurrence = occurrence[key]
        occurrence[key] += 1
        if match is not None:
            bank_amount = float(match.bank_amount)
            difference = round(bank_amount - float(match.gl_amount), 2)
            status = "Amount difference" if abs(difference) >= 0.005 else "Matched"
        else:
            bank_amount = None
            difference = None
            exception = gl_exceptions.get(row.row_id)
            status = "Duplicate" if exception and exception.category == "Possible duplicate" else "GL only"
        gl_rows.append(
            [row.date, row.account, row.description, row.reference, float(row.amount),
             duplicate_occurrence, bank_amount, difference, status]
        )
    return bank_rows, gl_rows


def _write_schedule(
    sheet,
    title_row: int,
    title: str,
    headers: list[str],
    rows: list[list[object]],
    amount_columns: tuple[int, ...],
) -> int:
    sheet.cell(title_row, 1, title)
    sheet.cell(title_row, 1).font = Font(name="Aptos", size=11, bold=True, color=DARK)
    header_row = title_row + 1
    for column, header in enumerate(headers, start=1):
        sheet.cell(header_row, column, header)
    _style_header(sheet, header_row, 1, len(headers))
    first_data_row = header_row + 1
    output_rows = rows or [[None] * len(headers)]
    for row_offset, values in enumerate(output_rows):
        row_number = first_data_row + row_offset
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, _excel_value(value))
            if column == 1 and value is not None:
                cell.number_format = DATE_FORMAT
            if column in amount_columns and value is not None:
                cell.number_format = ACCOUNTING
                cell.alignment = Alignment(horizontal="right", vertical="center")
    return first_data_row + len(output_rows)


def _write_bank_reconciliation(sheet, result: ReconciliationResult, period: str) -> None:
    gl_exceptions = result.exceptions[
        (result.exceptions["source"] == "GL")
        & (result.exceptions["category"] != "Possible duplicate")
    ].sort_values(["date", "source_row"])
    deposits = gl_exceptions[gl_exceptions["amount"] > 0]
    outstanding = gl_exceptions[gl_exceptions["amount"] < 0]

    sheet["A1"] = f"Bank Reconciliation Statement – {period}"
    sheet["A2"] = "Net movement basis (no independently confirmed opening or closing balances available)."
    for column, value in enumerate(["Item", None, None, "Amount", "Per schedule"], 1):
        sheet.cell(4, column, value)
    summary_rows = [
        (5, f"Net bank movement per statement ({period.split()[0]})", result.bank_balance, None),
        (6, "Add: deposits in transit (Sch. 1)", float(deposits["amount"].sum()), float(deposits["amount"].sum())),
        (7, "Less: outstanding cheques / payments (Sch. 2)", float(outstanding["amount"].sum()), float(outstanding["amount"].sum())),
        (8, "Adjusted bank movement", result.adjusted_bank_balance, None),
        (10, "Adjusted cash book movement (per Adjusted Cash Book)", result.adjusted_book_balance, None),
        (11, "Unreconciled difference", result.residual, None),
    ]
    for row, label, amount, schedule_amount in summary_rows:
        sheet.cell(row, 1, label)
        sheet.cell(row, 4, round(float(amount), 2)).number_format = ACCOUNTING
        if schedule_amount is not None:
            sheet.cell(row, 5, round(float(schedule_amount), 2)).number_format = ACCOUNTING

    deposit_rows = [[row.date, row.description, row.reference, float(row.amount)] for row in deposits.itertuples(index=False)]
    next_row = _write_schedule(
        sheet, 13, "Schedule 1 – Deposits in transit",
        ["Date", "Memo", "DocNo", "Amount"], deposit_rows, (4,)
    )
    outstanding_title_row = max(20, next_row + 2)
    outstanding_rows = [[row.date, row.description, row.reference, float(row.amount)] for row in outstanding.itertuples(index=False)]
    _write_schedule(
        sheet, outstanding_title_row, "Schedule 2 – Outstanding cheques / payments",
        ["Date", "Memo", "DocNo", "Amount"], outstanding_rows, (4,)
    )

    _apply_base_style(sheet)
    sheet["A1"].font = Font(name="Aptos", size=14, bold=True, color=DARK)
    sheet["A2"].font = Font(name="Aptos", size=11, italic=True, color=GREY)
    _style_header(sheet, 4, 1, 5)
    _style_header(sheet, 14, 1, 4)
    _style_header(sheet, outstanding_title_row + 1, 1, 4)
    _style_total(sheet, 8, 4)
    _style_total(sheet, 11, 4)
    sheet["D10"].font = Font(name="Aptos", size=11, color=GREEN)
    if not result.is_reconciled:
        sheet["D11"].fill = PatternFill("solid", fgColor=LIGHT_RED)
        sheet["D11"].font = Font(name="Aptos", size=11, bold=True, color=RED)
    sheet.conditional_formatting.add(
        "D11",
        CellIsRule(operator="notEqual", formula=["0"],
                   fill=PatternFill("solid", fgColor=LIGHT_RED),
                   font=Font(color=RED, bold=True)),
    )
    _set_widths(sheet, [14, 44, 18, 18, 18])
    sheet.freeze_panes = "A4"


def _write_adjusted_cash_book(sheet, result: ReconciliationResult, period: str) -> None:
    bank_only = result.exceptions[result.exceptions["source"] == "BANK"].sort_values(["date", "source_row"])
    duplicates = result.exceptions[
        (result.exceptions["source"] == "GL")
        & (result.exceptions["category"] == "Possible duplicate")
    ].sort_values(["date", "source_row"])
    variances = result.matches[result.matches["amount_variance"].abs() >= 0.005].sort_values(
        ["bank_date", "bank_source_row"]
    )

    sheet["A1"] = f"Adjusted Cash Book – {period}"
    sheet["A2"] = "Net movement basis. Bank-only items, duplicate GL postings and amount corrections are shown below."
    for column, value in enumerate(["Item", None, None, "Amount", "Per schedule", None], 1):
        sheet.cell(4, column, value)
    summary_rows = [
        (5, f"Net cash book movement per GL ({period.split()[0]})", result.gl_balance, None),
        (6, "Add/(less): bank items not recorded in cash book (Sch. A)", result.bank_only_adjustments, result.bank_only_adjustments),
        (7, "Add: reverse duplicate GL postings (Sch. B)", result.duplicate_reversals, result.duplicate_reversals),
        (8, "Add/(less): correct GL amount errors to bank (Sch. C)", result.matched_variance, result.matched_variance),
        (9, "Adjusted cash book movement", result.adjusted_book_balance, None),
    ]
    for row, label, amount, schedule_amount in summary_rows:
        sheet.cell(row, 1, label)
        sheet.cell(row, 4, round(float(amount), 2)).number_format = ACCOUNTING
        if schedule_amount is not None:
            sheet.cell(row, 5, round(float(schedule_amount), 2)).number_format = ACCOUNTING

    schedule_a_rows = [[row.date, row.description, row.reference, float(row.amount)] for row in bank_only.itertuples(index=False)]
    after_a = _write_schedule(
        sheet, 12, "Schedule A – Bank items not recorded in cash book",
        ["Date", "Description", "Reference", "Amount"], schedule_a_rows, (4,)
    )
    schedule_b_title = max(24, after_a + 3)
    schedule_b_rows = [[row.date, row.description, row.reference, -float(row.amount)] for row in duplicates.itertuples(index=False)]
    after_b = _write_schedule(
        sheet, schedule_b_title, "Schedule B – Duplicate GL postings to reverse",
        ["Date", "Memo", "DocNo", "Reversal"], schedule_b_rows, (4,)
    )
    schedule_c_title = max(30, after_b + 3)
    schedule_c_rows = [
        [row.bank_date, row.gl_description, row.gl_reference, float(row.gl_amount),
         float(row.bank_amount), float(row.amount_variance)]
        for row in variances.itertuples(index=False)
    ]
    _write_schedule(
        sheet, schedule_c_title, "Schedule C – GL amount errors (correct to bank amount)",
        ["Date", "Memo", "DocNo", "GL Amount", "Bank Amount", "Correction"],
        schedule_c_rows, (4, 5, 6)
    )

    _apply_base_style(sheet)
    sheet["A1"].font = Font(name="Aptos", size=14, bold=True, color=DARK)
    sheet["A2"].font = Font(name="Aptos", size=11, italic=True, color=GREY)
    _style_header(sheet, 4, 1, 5)
    _style_header(sheet, 13, 1, 4)
    _style_header(sheet, schedule_b_title + 1, 1, 4)
    _style_header(sheet, schedule_c_title + 1, 1, 6)
    _style_total(sheet, 9, 4)
    _set_widths(sheet, [14, 46, 18, 18, 18, 18])
    sheet.freeze_panes = "A4"


def build_excel_report(result: ReconciliationResult) -> bytes:
    """Return a four-tab workbook matching the accountant-facing reference layout."""

    period, suffix = _period_details(result)
    bank_rows, gl_rows = _source_views(result)

    workbook = Workbook()
    bank_sheet = workbook.active
    bank_sheet.title = f"bank_statement_{suffix}"[:31]
    gl_sheet = workbook.create_sheet(f"gl_cash_extract_{suffix}"[:31])
    reconciliation_sheet = workbook.create_sheet("Bank Reconciliation Statement")
    cash_book_sheet = workbook.create_sheet("Adjusted Cash Book")

    _write_source_sheet(
        bank_sheet,
        ["Date", "Description", "Reference", "Amount", "GL Count", "GL Amount", "Difference", "Status"],
        bank_rows, "BankReconciliationTable", [12, 42, 16, 15, 12, 15, 15, 20], (1,), (4, 6, 7)
    )
    _write_source_sheet(
        gl_sheet,
        ["Date", "Account", "Memo", "DocNo", "Amount", "Occurrence", "Bank Amount", "Difference", "Status"],
        gl_rows, "GLReconciliationTable", [12, 24, 44, 16, 15, 12, 15, 15, 20], (1,), (5, 7, 8)
    )
    _write_bank_reconciliation(reconciliation_sheet, result, period)
    _write_adjusted_cash_book(cash_book_sheet, result, period)

    workbook.active = 2
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
