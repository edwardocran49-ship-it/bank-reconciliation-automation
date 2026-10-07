from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from bank_recon import ReconciliationConfig, build_excel_report, reconcile


ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def sample_result():
    bank = pd.read_csv(ROOT / "data" / "raw" / "bank_statement_jun2026.csv")
    gl = pd.read_csv(ROOT / "data" / "raw" / "gl_cash_extract_jun2026.csv")
    return reconcile(bank, gl)


def test_sample_rows_are_fully_disposed_once(sample_result):
    result = sample_result
    assert len(result.bank) == 122
    assert len(result.gl) == 122
    assert result.matched_source_rows + len(result.exceptions) == result.total_source_rows
    assert result.matches["bank_row_id"].is_unique
    assert result.matches["gl_row_id"].is_unique


def test_sample_reconciliation_proof_balances(sample_result):
    assert sample_result.is_reconciled
    assert sample_result.residual == pytest.approx(0.0, abs=0.005)
    assert sample_result.adjusted_bank_balance == pytest.approx(
        sample_result.adjusted_book_balance, abs=0.005
    )
    assert sample_result.duplicate_reversals == pytest.approx(10_230.18)
    assert sample_result.bank_only_adjustments == pytest.approx(497.43)


def test_sample_exposes_multiple_control_categories(sample_result):
    categories = set(sample_result.exceptions["category"])
    assert "Possible duplicate" in categories
    assert "Outstanding check" in categories
    assert "Deposit in transit" in categories
    assert len(sample_result.matches) == 115
    assert len(sample_result.exceptions) == 14
    assert sample_result.matches["match_method"].value_counts().to_dict() == {
        "Same-day exact": 95,
        "Timing difference": 12,
        "Controlled variance": 8,
    }


def test_each_matching_pass_is_reachable():
    bank = pd.DataFrame(
        [
            ("2026-06-01", "Customer receipt", "A-1", 100.00),
            ("2026-06-05", "Supplier payment", "B-2", -75.00),
            ("2026-06-10", "Utility payment", "C-3", -50.40),
        ],
        columns=["Date", "Description", "Reference", "Amount"],
    )
    gl = pd.DataFrame(
        [
            ("2026-06-01", "Cash", "Customer receipt", "A-1", 100.00),
            ("2026-06-03", "Cash", "Supplier payment", "B-2", -75.00),
            ("2026-06-11", "Cash", "Utility payment", "C-3", -50.00),
        ],
        columns=["Date", "Account", "Memo", "DocNo", "Amount"],
    )

    result = reconcile(bank, gl, ReconciliationConfig(amount_tolerance=0.99))

    assert list(result.matches["match_method"]) == [
        "Same-day exact",
        "Timing difference",
        "Controlled variance",
    ]
    assert result.exceptions.empty
    assert result.residual == pytest.approx(0.0, abs=0.005)


def test_missing_required_column_is_rejected():
    bank = pd.DataFrame({"Date": ["2026-06-01"], "Amount": [10]})
    gl = pd.DataFrame(
        {
            "Date": ["2026-06-01"],
            "Account": ["Cash"],
            "Memo": ["Receipt"],
            "DocNo": ["A"],
            "Amount": [10],
        }
    )
    with pytest.raises(ValueError, match="missing required columns"):
        reconcile(bank, gl)


def test_excel_report_contains_workbook_bytes(sample_result):
    report = build_excel_report(sample_result)
    assert report[:2] == b"PK"
    assert len(report) > 10_000

    workbook = load_workbook(BytesIO(report), data_only=False)
    assert workbook.sheetnames == [
        "bank_statement_jun2026",
        "gl_cash_extract_jun2026",
        "Bank Reconciliation Statement",
        "Adjusted Cash Book",
    ]
    reconciliation = workbook["Bank Reconciliation Statement"]
    cash_book = workbook["Adjusted Cash Book"]
    assert reconciliation["D8"].value == pytest.approx(254_160.85)
    assert reconciliation["D11"].value == pytest.approx(0.0)
    assert cash_book["D7"].value == pytest.approx(10_230.18)
    assert cash_book["D9"].value == pytest.approx(254_160.85)
    assert all(
        not (isinstance(cell.value, str) and cell.value.startswith("#"))
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )
